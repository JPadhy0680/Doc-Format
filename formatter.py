from dataclasses import dataclass
from io import BytesIO
import re
from copy import deepcopy
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

@dataclass
class FormatOptions:
    document_type: str = "SOP"
    company_name: str = "Company Name"
    department: str = ""
    document_number: str = ""
    effective_date: str = ""
    review_date: str = ""
    add_toc: bool = False
    normalize_sections: bool = True
    add_end_marker: bool = True
    preserve_template_header_footer: bool = True

SECTION_MAP = {
 "revision summary":"Revision Summary", "amendment summary":"Amendment Summary",
 "objective":"Objective(s)", "objectives":"Objective(s)", "purpose":"Objective(s)",
 "scope":"Scope(s)", "scopes":"Scope(s)", "responsibility":"Responsibility(s)", "responsibilities":"Responsibility(s)",
 "accountability":"Accountability(s)", "accountabilities":"Accountability(s)",
 "abbreviation":"Abbreviation(s)", "abbreviations":"Abbreviation(s)",
 "definition":"Definition(s)", "definitions":"Definition(s)", "procedure":"Procedure(s)", "procedures":"Procedure(s)",
 "reference":"Reference(s)", "references":"Reference(s)", "list of annexure":"List of Annexure", "list of annexures":"List of Annexure",
 "table of contents":"Table of Contents", "overview":"Overview", "clause references":"Clause References", "applicability":"Applicability"
}
SOP_ORDER=["Revision Summary","Objective(s)","Scope(s)","Responsibility(s)","Accountability(s)","Abbreviation(s)","Definition(s)","Procedure(s)","Reference(s)","List of Annexure"]
IT_ORDER=["Amendment Summary","Table of Contents","Overview","Clause References","Purpose and Objectives","Scope","Applicability","Definitions","Procedure","List of Annexures"]

def _clean(text): return re.sub(r"\s+", " ", (text or "").strip())
def _key(text): return re.sub(r"[^a-z ]", "", _clean(text).lower()).strip()

def _set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tcPr=cell._tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in("w:tcMar")
    if tcMar is None: tcMar=OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=tcMar.find(qn('w:'+m))
        if node is None: node=OxmlElement('w:'+m); tcMar.append(node)
        node.set(qn('w:w'),str(v)); node.set(qn('w:type'),'dxa')

def _field(run, instruction):
    begin=OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'),'begin')
    instr=OxmlElement('w:instrText'); instr.set(qn('xml:space'),'preserve'); instr.text=instruction
    sep=OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'),'separate')
    end=OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'),'end')
    run._r.extend([begin,instr,sep,end])

def _setup_styles(doc):
    normal=doc.styles['Normal']; normal.font.name='Times New Roman'; normal.font.size=Pt(12)
    normal._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman')
    pf=normal.paragraph_format; pf.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY; pf.line_spacing=1.5; pf.space_after=Pt(6)
    for level,size in [(1,12),(2,12),(3,12)]:
        s=doc.styles[f'Heading {level}']; s.font.name='Times New Roman'; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor(0,0,0)
        s._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman')
        s.paragraph_format.space_before=Pt(10); s.paragraph_format.space_after=Pt(4); s.paragraph_format.keep_with_next=True
    if 'Procedure Level 1' not in [s.name for s in doc.styles]:
        for name,left,first in [('Procedure Level 1',0,0),('Procedure Level 2',0.25,-0.25),('Procedure Level 3',0.5,-0.25)]:
            s=doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH); s.base_style=doc.styles['Normal']; s.font.bold=(name=='Procedure Level 1')
            s.paragraph_format.left_indent=Inches(left); s.paragraph_format.first_line_indent=Inches(first); s.paragraph_format.keep_with_next=True

def _copy_body(src, dst):
    body=dst.element.body
    sect=body.sectPr
    for child in list(body):
        if child is not sect: body.remove(child)
    for child in src.element.body:
        if child.tag != qn('w:sectPr'): body.insert(len(body)-1, deepcopy(child))

def _header_footer(doc,opt):
    sec=doc.sections[0]; sec.different_first_page_header_footer=True
    if not opt.preserve_template_header_footer or not sec.header.paragraphs[0].text.strip():
        h=sec.header; p=h.paragraphs[0]; p.text=''
        t=h.add_table(rows=2, cols=3, width=Inches(6.5)); t.alignment=WD_TABLE_ALIGNMENT.CENTER
        t.cell(0,0).text=opt.company_name or 'Company Name'; t.cell(0,1).text=opt.document_type; t.cell(0,2).text=opt.document_number or 'Document No.'
        t.cell(1,0).text=opt.department or 'Department'; t.cell(1,1).text=f"Effective: {opt.effective_date or 'DD-MMM-YYYY'}"; t.cell(1,2).text=f"Review: {opt.review_date or 'DD-MMM-YYYY'}"
        for row in t.rows:
            for c in row.cells:
                _set_cell_margins(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for p in c.paragraphs:
                    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs: r.font.name='Times New Roman'; r.font.size=Pt(9); r.bold=True
    if not opt.preserve_template_header_footer or not sec.footer.paragraphs[0].text.strip():
        p=sec.footer.paragraphs[0]; p.text=''; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.add_run('Page '); _field(p.add_run(),'PAGE'); p.add_run(' of '); _field(p.add_run(),'NUMPAGES')
        for r in p.runs: r.font.name='Times New Roman'; r.font.size=Pt(9)

def _classify_and_format(doc,opt,report):
    section_no=0; in_procedure=False; recognized=[]
    for p in doc.paragraphs:
        txt=_clean(p.text)
        if not txt: continue
        k=_key(txt.rstrip(':'))
        if k in SECTION_MAP:
            section_no += 1; title=SECTION_MAP[k]
            if opt.normalize_sections: p.text=f"{_roman(section_no)}. {title}"
            p.style='Heading 1'; in_procedure=title.startswith('Procedure'); recognized.append(title); continue
        if txt.upper()=='END OF DOCUMENT':
            p.style='Heading 1'; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; in_procedure=False; continue
        m=re.match(r'^((?:\d+\.)+\d*|\d+\.0)\s+(.+)$',txt)
        if in_procedure and m:
            depth=min(3,m.group(1).count('.')+1 if not m.group(1).endswith('.0') else 1)
            p.style=f'Procedure Level {depth}'; continue
        if re.match(r'^[\-•▪]\s+',txt):
            p.text=re.sub(r'^[\-•▪]\s+','',txt); p.style='List Bullet'; continue
        if len(txt)<=80 and (txt.endswith(':') or (txt.isupper() and len(txt.split())<=8)):
            p.style='Heading 2'; continue
        p.style='Normal'; p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    report.append(('INFO',f'Recognized sections: {", ".join(recognized) if recognized else "none"}.'))

def _roman(n):
    vals=[(10,'X'),(9,'IX'),(5,'V'),(4,'IV'),(1,'I')]; out=''
    for v,s in vals:
        while n>=v: out+=s; n-=v
    return out

def _format_tables(doc):
    for table in doc.tables:
        table.alignment=WD_TABLE_ALIGNMENT.CENTER
        try: table.style='Table Grid'
        except KeyError: pass
        for i,row in enumerate(table.rows):
            row._tr.get_or_add_trPr()
            for c in row.cells:
                _set_cell_margins(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for p in c.paragraphs:
                    p.paragraph_format.line_spacing=1.0; p.paragraph_format.space_after=Pt(0)
                    for r in p.runs: r.font.name='Times New Roman'; r.font.size=Pt(10); r.bold=(i==0)
        if table.rows:
            trPr=table.rows[0]._tr.get_or_add_trPr(); rep=OxmlElement('w:tblHeader'); rep.set(qn('w:val'),'true'); trPr.append(rep)

def _insert_toc(doc):
    p=doc.paragraphs[0].insert_paragraph_before('Table of Contents'); p.style='Heading 1'
    toc=p.insert_paragraph_before(); run=toc.add_run(); _field(run,'TOC \\o "1-3" \\h \\z \\u')

def _add_end(doc):
    if not any(_clean(p.text).upper()=='END OF DOCUMENT' for p in doc.paragraphs):
        p=doc.add_paragraph('END OF DOCUMENT',style='Heading 1'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER

def format_document(source_bytes, template_bytes, options):
    report=[]
    src=Document(BytesIO(source_bytes))
    if template_bytes:
        out=Document(BytesIO(template_bytes)); _copy_body(src,out); report.append(('PASS','Approved template loaded; source body content copied into it.'))
    else:
        out=src; report.append(('WARNING','No approved template supplied. A generic controlled-document layout was applied.'))
    _setup_styles(out); _header_footer(out,options); _classify_and_format(out,options,report); _format_tables(out)
    if options.add_toc: _insert_toc(out); report.append(('INFO','An updateable Word TOC field was inserted. In Word, right-click it and select Update Field.'))
    if options.add_end_marker: _add_end(out)
    bio=BytesIO(); out.save(bio); return bio.getvalue(),report
