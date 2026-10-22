from io import BytesIO
import re
from docx import Document

def validate_document(data, options):
    doc=Document(BytesIO(data)); text='\n'.join(p.text.strip() for p in doc.paragraphs if p.text.strip()); low=text.lower(); r=[]
    required = ['objective','scope','responsibil','accountabil','procedure','list of annexure'] if options.document_type!='IT Policy' else ['amendment summary','overview','scope','applicability','definition','procedure','list of annexure']
    for item in required:
        r.append(('PASS',f'Required section detected: {item}.') if item in low else ('WARNING',f'Required section was not confidently detected: {item}.'))
    r.append(('PASS','END OF DOCUMENT marker is present.') if 'end of document' in low else ('ERROR','END OF DOCUMENT marker is missing.'))
    if options.document_number:
        pattern=r'^[A-Za-z0-9]+/(SOP|WG|Policy)/[A-Za-z]{2,4}/\d{2}-\d{2}$'
        r.append(('PASS','Document number follows the configured pattern.') if re.match(pattern,options.document_number,re.I) else ('WARNING','Document number may not follow XXX/TYPE/DEPT/YY-00.'))
    else: r.append(('WARNING','Document number was not provided.'))
    date_re=r'^\d{2}-[A-Za-z]{3}-\d{4}$'
    for name,val in [('Effective date',options.effective_date),('Review date',options.review_date)]:
        if not val: r.append(('WARNING',f'{name} was not provided.'))
        elif not re.match(date_re,val): r.append(('WARNING',f'{name} is not in DD-MMM-YYYY format.'))
        else: r.append(('PASS',f'{name} format is valid.'))
    if not doc.tables: r.append(('INFO','No source tables were detected.'))
    else: r.append(('PASS',f'{len(doc.tables)} table(s) formatted.'))
    return r
