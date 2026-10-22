from io import BytesIO
from docx import Document
from formatter import FormatOptions, format_document

def test_format_smoke():
    d=Document(); d.add_paragraph('PURPOSE'); d.add_paragraph('Test body.'); d.add_paragraph('PROCEDURE'); d.add_paragraph('1.0 First step')
    b=BytesIO(); d.save(b)
    out, report=format_document(b.getvalue(),None,FormatOptions())
    result=Document(BytesIO(out))
    assert any(p.text=='END OF DOCUMENT' for p in result.paragraphs)
    assert len(out)>0
